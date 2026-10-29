import io

from app.services.extraction.enums import ChunkType, ExtractorType
from app.services.extraction.types import ExtractedBlock, ExtractionPayload

DOCX_VERSION = "docx-python-docx-v1"


def extract_docx(data: bytes) -> ExtractionPayload:
    from docx import Document
    from docx.oxml.ns import qn

    document = Document(io.BytesIO(data))
    heading_path: list[str] = []
    blocks: list[ExtractedBlock] = []
    title: str | None = None
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        style = paragraph.style.name if paragraph.style is not None else ""
        if style.startswith("Heading") and text:
            level = _heading_level(style)
            if title is None:
                title = text
            heading_path = heading_path[: level - 1]
            heading_path.append(text)
            blocks.append(ExtractedBlock(ChunkType.SECTION.value, text, list(heading_path), {"level": level}))
            continue
        if text:
            blocks.append(ExtractedBlock(ChunkType.PARAGRAPH.value, text, list(heading_path), {}))
    for table in document.tables:
        rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(" | ".join(cells))
        if rows:
            blocks.append(ExtractedBlock(ChunkType.TABLE.value, "\n".join(rows), list(heading_path), {}))
    image_count = 0
    for element in document.element.iter(qn("a:blip")):
        image_count += 1
    return ExtractionPayload(
        extractor_type=ExtractorType.DOCX.value,
        extractor_version=DOCX_VERSION,
        title=title,
        blocks=blocks,
        metadata={"embedded_image_count": image_count},
    )


def _heading_level(style_name: str) -> int:
    digits = "".join(char for char in style_name if char.isdigit())
    if not digits:
        return 1
    return max(1, min(int(digits), 6))
