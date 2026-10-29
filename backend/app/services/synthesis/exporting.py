import os
import re
from io import BytesIO
from pathlib import Path

import fitz
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.core.config import PROJECT_ROOT


def safe_filename(value: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*\s]+', "_", value).strip("._")
    return cleaned[:80] or "brief"


def markdown_name(brief) -> str:
    categories = "_".join(brief.filters_json.get("category_slugs") or ["all"])
    start = str(brief.start_date)
    end = str(brief.end_date)
    span = start if start == end else f"{start}_{end}"
    return safe_filename(f"{span}_{categories}_学习摘要") + ".md"


def write_markdown(brief) -> bytes:
    return (brief.content_markdown or "").encode("utf-8")


def write_docx(brief, sources: list[dict]) -> bytes:
    document = Document()
    document.add_heading(brief.title or "学习摘要", level=1)
    for line in (brief.content_markdown or "").splitlines():
        if line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            continue
        elif line.startswith("URL: "):
            url = line.removeprefix("URL: ").strip()
            paragraph = document.add_paragraph("")
            _hyperlink(paragraph, url, url)
        else:
            document.add_paragraph(line)
    for source in sources:
        url = source.get("canonical_url") or source.get("original_url")
        if url:
            paragraph = document.add_paragraph(f"[{source['citation_number']}] ")
            _hyperlink(paragraph, source["title"], url)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def write_pdf(brief) -> bytes:
    document = fitz.open()
    font_path = _cjk_font()
    text = brief.content_markdown or ""
    chunks = [text[index : index + 1200] for index in range(0, max(len(text), 1), 1200)]
    for chunk in chunks:
        page = document.new_page()
        if font_path is not None:
            page.insert_font(fontname="cjk", fontfile=str(font_path))
            page.insert_textbox(fitz.Rect(48, 48, 548, 800), chunk, fontname="cjk", fontsize=11)
        else:
            page.insert_textbox(fitz.Rect(48, 48, 548, 800), chunk, fontsize=11)
    payload = document.tobytes()
    document.close()
    return payload


def cache_export(name: str, payload: bytes) -> Path:
    root = PROJECT_ROOT / "data" / "exports"
    root.mkdir(parents=True, exist_ok=True)
    path = root / safe_filename(name)
    temporary = path.with_name(path.name + ".partial")
    temporary.write_bytes(payload)
    os.replace(temporary, path)
    return path


def _cjk_font() -> Path | None:
    candidates = [Path(r"C:\Windows\Fonts\msyh.ttc"), Path(r"C:\Windows\Fonts\simsun.ttc"), Path(r"C:\Windows\Fonts\simhei.ttf")]
    return next((path for path in candidates if path.exists()), None)


def _hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    run_properties.append(color)
    run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
