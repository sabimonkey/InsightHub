import io
from pathlib import Path

import fitz
import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw, ImageFont
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.api.deps import get_extraction_service, get_ingestion_service
from app.main import app
from app.models.extraction import ExtractedDocument, ExtractionRun
from app.models.ingestion_record import IngestionRecord
from app.services.extraction.chunking import chunk_blocks
from app.services.extraction.enums import ChunkType
from app.services.extraction.service import ExtractionService
from app.services.extraction.ssrf import SsrfError, assert_public_http_url
from app.services.extraction.types import ExtractedBlock
from app.services.extraction.web_extractor import extract_web
from app.services.ingestion.service import IngestionService


def _service(session: Session, raw_root: Path, tmp_path: Path, **kwargs) -> ExtractionService:
    return ExtractionService(
        session,
        raw_root,
        tmp_path / "extracted",
        inline=True,
        chunk_max_chars=kwargs.get("chunk_max_chars", 2000),
        chunk_overlap_chars=kwargs.get("chunk_overlap_chars", 50),
        max_pdf_pages=kwargs.get("max_pdf_pages", 5),
        web_timeout=2,
        web_max_bytes=kwargs.get("web_max_bytes", 100_000),
        web_max_redirects=kwargs.get("web_max_redirects", 2),
        ocr_image=kwargs.get("ocr_image"),
        http_client=kwargs.get("http_client"),
        resolver=kwargs.get("resolver"),
    )


def test_text_extraction_creates_document_and_chunk(client, db_session):
    created = client.post("/api/ingest/text", json={"text": "Hello extraction\nsecond line"})
    source_id = created.json()["source"]["id"]
    service = _service(db_session, Path("."), Path("."))
    from app.services.extraction.enums import ExtractionTrigger

    run = service.schedule(source_id, ExtractionTrigger.MANUAL.value, force=True)
    assert run.status == "SUCCEEDED"
    view_run, document, chunks = service.current_view(source_id)
    assert view_run.status == "SUCCEEDED"
    assert document is not None
    assert "Hello extraction" in document.plain_text
    assert "second line" in document.plain_text
    assert document.structured_text is None
    assert chunks
    assert chunks[0].chunk_type == "PARAGRAPH"
    assert "Hello extraction" in chunks[0].content


def test_markdown_keeps_heading_path(db_session, raw_root, tmp_path):
    markdown = "# Agent\n\nIntro\n\n## Memory\n\nRemember this\n\n```python\nprint(1)\n```\n"
    path = tmp_path / "note.md"
    path.write_text(markdown, encoding="utf-8")
    from app.services.extraction.markdown_extractor import extract_markdown

    payload = extract_markdown(path.read_text(encoding="utf-8"))
    paths = [block.heading_path for block in payload.blocks if block.chunk_type == "PARAGRAPH"]
    assert ["Agent"] in paths
    assert ["Agent", "Memory"] in paths
    assert any(block.chunk_type == "CODE" and "print(1)" in block.content for block in payload.blocks)
    service = _service(db_session, raw_root, tmp_path, chunk_max_chars=30, chunk_overlap_chars=5)
    huge = ExtractedBlock(ChunkType.PARAGRAPH.value, "A" * 80, ["Agent", "Memory"], {})
    pieces = chunk_blocks([huge], service.chunk_max_chars, service.chunk_overlap_chars)
    assert len(pieces) > 1
    assert all(piece.heading_path == ["Agent", "Memory"] for piece in pieces)


def test_docx_extracts_paragraph_and_table(tmp_path):
    document = Document()
    document.add_heading("Report", level=1)
    document.add_paragraph("Body text")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Alpha"
    table.cell(1, 1).text = "1"
    target = tmp_path / "sample.docx"
    document.save(target)
    from app.services.extraction.docx_extractor import extract_docx

    payload = extract_docx(target.read_bytes())
    assert any("Body text" in block.content for block in payload.blocks)
    table_blocks = [block for block in payload.blocks if block.chunk_type == "TABLE"]
    assert table_blocks
    assert "Alpha" in table_blocks[0].content
    assert payload.metadata["embedded_image_count"] == 0


def test_scanned_pdf_records_ocr_method(tmp_path):
    pdf = fitz.open()
    pdf.new_page()
    buffer = io.BytesIO()
    pdf.save(buffer)
    pdf.close()
    from app.services.extraction.pdf_extractor import extract_pdf

    payload = extract_pdf(
        buffer.getvalue(),
        max_pages=5,
        ocr_image=lambda _data: "SCANNED WORD",
        derived_dir=tmp_path / "derived",
    )
    assert payload.extractor_type == "PDF_OCR"
    assert payload.blocks[0].metadata["page_number"] == 1
    assert payload.blocks[0].metadata["method"] == "ocr"
    assert "SCANNED WORD" in payload.blocks[0].content
    assert (tmp_path / "derived" / "page-1.png").is_file()
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Visible PDF sentence")
    buffer = io.BytesIO()
    pdf.save(buffer)
    pdf.close()
    from app.services.extraction.pdf_extractor import extract_pdf

    payload = extract_pdf(buffer.getvalue(), max_pages=5, ocr_image=None, derived_dir=tmp_path / "derived")
    assert payload.extractor_type == "PDF_TEXT"
    assert payload.status == "SUCCEEDED"
    assert "Visible PDF sentence" in payload.blocks[0].content
    assert payload.blocks[0].metadata["page_number"] == 1


def test_corrupt_pdf_fails_without_changing_ingestion(client, db_session, raw_root):
    response = client.post("/api/ingest/file", files={"file": ("broken.pdf", b"%PDF-1.4\nbad", "application/pdf")})
    assert response.status_code == 200
    body = response.json()
    assert body["ingestion"]["status"] == "STORED"
    stored = raw_root / body["asset"]["stored_path"]
    before = stored.read_bytes()
    from app.services.extraction.enums import ExtractionTrigger

    service = _service(db_session, raw_root, raw_root.parent)
    run = service.schedule(body["source"]["id"], ExtractionTrigger.MANUAL.value, force=True)
    assert run.status == "FAILED"
    assert stored.read_bytes() == before
    record = db_session.get(IngestionRecord, body["ingestion"]["id"])
    assert record.status == "STORED"


def test_duplicate_collection_does_not_extract_again(db_session, raw_root, tmp_path):
    from app.services.extraction.enums import ExtractionTrigger

    service = _service(db_session, raw_root, tmp_path)
    ingest = IngestionService(db_session, raw_root, max_upload_size_bytes=10_000)
    first = ingest.ingest_text("same text body")
    run = service.schedule(first.source_item.id, ExtractionTrigger.AUTO.value, force=False)
    assert run.status == "SUCCEEDED"
    second = ingest.ingest_text("same text body")
    assert second.duplicate_detected is True
    again = service.schedule(second.source_item.id, ExtractionTrigger.AUTO.value, force=False)
    assert again.id == run.id
    count = db_session.scalar(select(func.count()).select_from(ExtractionRun).where(ExtractionRun.source_item_id == first.source_item.id))
    assert count == 1


def test_web_html_extracts_og_image_metadata():
    html = """
    <html><head><meta property="og:image" content="https://example.com/hero.jpg" /></head>
    <body><p>Article body with enough visible text for extraction path.</p></body></html>
    """
    from app.services.extraction.web_content import payload_from_html

    payload = payload_from_html(html, {"requested_url": "https://example.com/post", "final_url": "https://example.com/post"})
    assert payload.metadata.get("preview_image_url") == "https://example.com/hero.jpg"


def test_ssrf_blocks_local_and_redirect(tmp_path):
    with pytest.raises(SsrfError):
        assert_public_http_url("http://127.0.0.1/secret")
    with pytest.raises(SsrfError):
        assert_public_http_url("http://10.1.2.3/secret")

    def public_resolver(_host: str) -> list[str]:
        return ["93.184.216.34"]

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/start":
            return httpx.Response(302, headers={"location": "http://127.0.0.1/private"})
        return httpx.Response(200, text="<html><title>Ok</title><body><p>Public article text that is long enough.</p></body></html>")

    transport = httpx.MockTransport(handler)
    client = httpx.Client(transport=transport, follow_redirects=False)
    blocked = extract_web(
        "http://example.test/start",
        timeout=2,
        max_bytes=10000,
        max_redirects=3,
        client=client,
        resolver=public_resolver,
    )
    assert blocked.status == "FAILED"
    assert blocked.error_code == "SSRF_BLOCKED"
    allowed = extract_web(
        "http://example.test/page",
        timeout=2,
        max_bytes=10000,
        max_redirects=3,
        client=client,
        resolver=public_resolver,
    )
    assert allowed.status == "SUCCEEDED"
    assert allowed.title == "Ok"
    assert allowed.metadata["http_status"] == 200
    assert "Public article text" in allowed.blocks[0].content


def test_reextract_keeps_history(db_session, raw_root, tmp_path):
    from app.services.extraction.enums import ExtractionTrigger

    ingest = IngestionService(db_session, raw_root, max_upload_size_bytes=10_000)
    created = ingest.ingest_text("versioned text")
    service = _service(db_session, raw_root, tmp_path)
    first = service.schedule(created.source_item.id, ExtractionTrigger.MANUAL.value, force=True)
    second = service.schedule(created.source_item.id, ExtractionTrigger.REEXTRACT.value, force=True)
    assert first.id != second.id
    assert second.is_current is True
    count = db_session.scalar(select(func.count()).select_from(ExtractionRun).where(ExtractionRun.source_item_id == created.source_item.id))
    assert count == 2


def test_image_ocr_finds_key_text(tmp_path):
    image = Image.new("RGB", (640, 200), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=48)
    draw.text((40, 70), "HELLO", fill="black", font=font)
    target = tmp_path / "hello.png"
    image.save(target)
    from app.services.extraction.image_extractor import extract_image, recognize_image_bytes

    payload = extract_image(target.read_bytes(), ".png", recognize_image_bytes)
    assert payload.status == "SUCCEEDED"
    assert payload.extractor_type == "IMAGE_OCR"
    assert "HELLO" in payload.blocks[0].content.upper()


class _CountingStream(httpx.SyncByteStream):
    def __init__(self, pieces: list[bytes], counter: dict[str, int]) -> None:
        self.pieces = pieces
        self.counter = counter

    def __iter__(self):
        for piece in self.pieces:
            self.counter["chunks"] += 1
            yield piece


def _public_resolver(_host: str) -> list[str]:
    return ["93.184.216.34"]


def test_web_stream_stops_without_content_length(tmp_path):
    counter = {"chunks": 0}
    pieces = [b"A" * 1000 for _ in range(20)]

    def handler(_request: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200,
            stream=_CountingStream(pieces, counter),
            headers={"content-type": "text/html"},
        )

    client = httpx.Client(transport=httpx.MockTransport(handler), follow_redirects=False)
    payload = extract_web(
        "http://example.test/huge",
        timeout=2,
        max_bytes=2500,
        max_redirects=2,
        client=client,
        resolver=_public_resolver,
    )
    assert payload.status == "FAILED"
    assert payload.error_code == "WEB_RESPONSE_TOO_LARGE"
    assert payload.blocks == []
    assert 0 < counter["chunks"] < len(pieces)


def test_web_rejects_declared_content_length_without_reading_body():
    counter = {"chunks": 0}

    def handler(_request: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200,
            stream=_CountingStream([b"x" * 100], counter),
            headers={"content-type": "text/html", "content-length": "999999"},
        )

    client = httpx.Client(transport=httpx.MockTransport(handler), follow_redirects=False)
    payload = extract_web(
        "http://example.test/declared",
        timeout=2,
        max_bytes=1000,
        max_redirects=1,
        client=client,
        resolver=_public_resolver,
    )
    assert payload.status == "FAILED"
    assert payload.error_code == "WEB_RESPONSE_TOO_LARGE"
    assert counter["chunks"] == 0


def test_web_small_response_still_succeeds():
    html = "<html><title>Small</title><body><p>Public article text that is long enough to keep.</p></body></html>"

    def handler(_request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, text=html, headers={"content-type": "text/html", "content-length": str(len(html))})

    client = httpx.Client(transport=httpx.MockTransport(handler), follow_redirects=False)
    payload = extract_web(
        "http://example.test/small",
        timeout=2,
        max_bytes=10_000,
        max_redirects=1,
        client=client,
        resolver=_public_resolver,
    )
    assert payload.status == "SUCCEEDED"
    assert payload.error_code is None
    assert "Public article text" in payload.blocks[0].content


def test_oversized_web_response_does_not_save_document(db_session, raw_root, tmp_path):
    counter = {"chunks": 0}

    def handler(_request: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200,
            stream=_CountingStream([b"B" * 400 for _ in range(10)], counter),
            headers={"content-type": "text/html"},
        )

    client = httpx.Client(transport=httpx.MockTransport(handler), follow_redirects=False)
    ingest = IngestionService(db_session, raw_root, max_upload_size_bytes=10_000)
    created = ingest.ingest_url("http://example.test/over-limit")
    from app.services.extraction.enums import ExtractionTrigger

    service = _service(
        db_session,
        raw_root,
        tmp_path,
        web_max_bytes=500,
        http_client=client,
        resolver=_public_resolver,
    )
    run = service.schedule(created.source_item.id, ExtractionTrigger.MANUAL.value, force=True)
    assert run.status == "FAILED"
    assert run.error_code == "WEB_RESPONSE_TOO_LARGE"
    _view_run, document, chunks = service.current_view(created.source_item.id)
    assert document is None
    assert chunks == []
    stored = db_session.scalar(select(func.count()).select_from(ExtractedDocument).where(ExtractedDocument.source_item_id == created.source_item.id))
    assert stored == 0
    assert counter["chunks"] < 10


def test_real_scanned_pdf_ocr_recovers_text(db_session, raw_root, tmp_path):
    image = Image.new("RGB", (980, 220), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=52)
    draw.text((30, 70), "HELLO PDF OCR", fill="black", font=font)
    png = io.BytesIO()
    image.save(png, format="PNG")
    pdf = fitz.open()
    page = pdf.new_page(width=980, height=220)
    page.insert_image(page.rect, stream=png.getvalue())
    assert len(page.get_text("text").strip()) < 20
    buffer = io.BytesIO()
    pdf.save(buffer)
    pdf.close()
    ingest = IngestionService(db_session, raw_root, max_upload_size_bytes=2_000_000)
    created = ingest.ingest_file("scan.pdf", io.BytesIO(buffer.getvalue()), "application/pdf")
    from app.services.extraction.enums import ExtractionTrigger

    service = _service(db_session, raw_root, tmp_path)
    run = service.schedule(created.source_item.id, ExtractionTrigger.MANUAL.value, force=True)
    assert run.status == "SUCCEEDED"
    assert run.extractor_type == "PDF_OCR"
    _view_run, document, chunks = service.current_view(created.source_item.id)
    assert document is not None
    assert document.extractor_type == "PDF_OCR"
    assert "HELLO" in document.plain_text.upper()
    assert chunks
    assert chunks[0].metadata_json["page_number"] == 1
    assert chunks[0].metadata_json["method"] == "ocr"


def test_extract_api_and_failure_status(db_session, raw_root, tmp_path):
    def override_ingestion():
        yield IngestionService(db_session, raw_root, max_upload_size_bytes=50_000)

    def override_extraction():
        yield _service(db_session, raw_root, tmp_path)

    app.dependency_overrides[get_ingestion_service] = override_ingestion
    app.dependency_overrides[get_extraction_service] = override_extraction
    try:
        with TestClient(app) as test_client:
            created = test_client.post("/api/ingest/text", json={"text": "API extracted body"})
            source_id = created.json()["source"]["id"]
            detail = test_client.get(f"/api/sources/{source_id}")
            assert detail.status_code == 200
            body = detail.json()
            assert body["extraction"]["run"]["status"] == "SUCCEEDED"
            assert "API extracted body" in body["extraction"]["document"]["plain_text"]
            again = test_client.post(f"/api/sources/{source_id}/reextract")
            assert again.status_code == 200
            assert again.json()["run"]["trigger"] == "REEXTRACT"
            inbox = test_client.get("/api/inbox")
            assert inbox.json()["items"][0]["extraction_status"] == "SUCCEEDED"
    finally:
        app.dependency_overrides.clear()
